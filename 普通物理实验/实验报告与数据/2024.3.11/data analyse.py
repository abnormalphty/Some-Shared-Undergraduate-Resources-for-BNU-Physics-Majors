import numpy as np
import matplotlib.pyplot as plt
import scipy.optimize as op
import openpyxl as xl
import os

data1 = []
data2 = []
data3 = []
count = 0
#数据读取
text = open(r'D:\document\for myself\2024.2-2024.6\普物实验2\2024.3.11\1092.txt','r',encoding='utf-8')
while True:
    file = text.readline()
    if not file:
        break
    count+=1
    if count%1==0:
        data = file.split()
        data1.append(data[0])
        data2.append(data[1])
        data3.append(data[2])
data1 = np.array(data1)
data2 = np.array(data2)
data3 = np.array(data3)
t_data = np.linspace(0,161.99,16200)
print(count)
#定义拟合函数
def func1(x,alpha,A_1,A_2,f_1,f_2,phi_1,phi_2,C):
    return np.exp(-alpha*x)*(A_1*np.sin(f_1*x+phi_1)+A_2*np.sin(f_2*x+phi_2))+C
def func3(x,A,w,phi,C):
    return A*np.sin(w*x+phi)+C
#拟合
#popt, pcov = curve_fit(func1,t_data,data1)
#print(popt)

#plt.scatter(t_data,data1)
#plt.plot(t_data,func1(t_data,*popt),'ro')
#plt.legend()
#plt.show()

#data9=xl.Workbook()
#sheet1=data9.create_sheet('all')
#for i in range(1,16291):
    #sheet1.cell(row=i,column=1,value=0.01*i-0.01)
    #sheet1.cell(row=i,column=2,value=data1[i-1])
#data9.save(r'D:\document\for myself\2024.2-2024.6\普物实验2\2024.3.11\JData.xlsx')

plt.grid()
plt.xlabel('time/s')
plt.ylabel('U_C/V')
plt.xlim(0,150)
plt.ylim(-20,20)
plt.title('RLC')

plt.plot(t_data,data1,'b-')


param_bounds=([0,2.8,-np.inf,3.37,0.13,-np.pi,-np.pi,-np.inf],[np.inf,np.inf,np.inf,3.41,0.15,np.pi,np.pi,np.inf])
popt,pcov=op.curve_fit(func1,t_data,data1,bounds=param_bounds)
print(popt)
plt.plot(t_data,func1(t_data,*popt),'r--')
plt.show()

