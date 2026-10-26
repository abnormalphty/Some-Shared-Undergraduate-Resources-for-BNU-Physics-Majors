from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


OUT = Path(__file__).resolve().parent / "Part3教学设计.docx"


def set_run(run, size=11, bold=False, color="000000"):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}" if level <= 3 else "Normal"
    run = p.add_run(text)
    set_run(run, size=16 if level == 1 else 13, bold=True, color="C00000" if level == 1 else "44546A")
    return p


def add_para(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run(run)
    return p


def add_item(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-18)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    r1 = p.add_run(label)
    set_run(r1, bold=True)
    r2 = p.add_run(text)
    set_run(r2)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, center=False, size=9):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run(run, size=size, bold=bold)


def add_evaluation_table(doc):
    headers = ["评价项目", "评价标准", "自评", "组评", "师评"]
    rows = [
        ["作品效果\n（30分）", "能稳定实现上浮、下沉和悬浮：30分；能实现上浮和下沉：20分；只能完成一种状态：10分。", "____", "____", "____"],
        ["原理解释\n（25分）", "能用重力与浮力的大小关系解释三种状态，并说明挤压后水进入浮沉子、总重力增大。", "____", "____", "____"],
        ["制作调试\n（20分）", "操作规范；能逐次微调配重或气室；如实记录现象、问题、调整方法和结果。", "____", "____", "____"],
        ["合作表现\n（15分）", "分工明确，人人参与；能倾听同伴意见，共同分析故障并完成改进。", "____", "____", "____"],
        ["展示反思\n（10分）", "表达清楚；能说出一次失败原因、解决办法及下一步改进设想。", "____", "____", "____"],
        ["总分", "满分100分", "____", "____", "____"],
    ]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, "44546A")
        set_cell_text(cell, text, bold=True, center=True, size=9)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            if i == 0:
                shade_cell(cells[i], "F2F2F2")
            set_cell_text(cells[i], text, bold=(i == 0), center=(i != 1), size=8.5)
    return table


def setup(doc):
    sec = doc.sections[0]
    sec.top_margin = Pt(72)
    sec.bottom_margin = Pt(72)
    sec.left_margin = Pt(72)
    sec.right_margin = Pt(72)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(11)


def build():
    doc = Document()
    setup(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Part3 教学设计")
    set_run(r, size=20, bold=True, color="C00000")

    add_heading(doc, "一、教学设计思路", 1)
    add_para(
        doc,
        "本节课围绕“制作浮沉子”这一跨学科实践活动展开，采用项目式学习的方式组织教学。课堂以“让浮沉子听指挥，实现上浮、下沉和悬浮”为任务驱动，引导学生在观察现象、分析原理、动手制作、测试调试和展示交流的过程中，理解浮沉条件、压强传递以及重力变化对物体浮沉状态的影响。",
    )
    add_para(
        doc,
        "整个教学设计坚持“做中学、悟中学、评中学”的思路。学生不是被动接受浮力知识，而是在真实制作任务中不断发现问题、提出猜想、验证原因、改进方案，从而把物理概念、工程思维和合作探究融合起来。",
    )

    add_heading(doc, "二、教学过程设计", 1)
    add_item(
        doc,
        "1. 情境导入，提出任务：",
        "教师演示浮沉子在水中随手控制上浮、下沉和悬浮的现象，激发学生好奇心，并提出本节课的项目任务：制作一个能“听指挥”的小潜水员。",
    )
    add_item(
        doc,
        "2. 观察现象，提出问题：",
        "学生观察浮沉子在自然状态和挤压瓶身后的变化，思考“为什么一挤瓶子它就会下沉”“松手后为什么又会上浮”，初步形成关于浮沉原因的猜想。",
    )
    add_item(
        doc,
        "3. 分析原理，突破迷思：",
        "教师引导学生从受力角度分析浮沉条件，明确浮沉子的浮沉取决于重力和浮力的大小关系。重点追问“进入小瓶的水改变了谁”，帮助学生认识到浮沉子下沉的主要原因是自身总重力增大，而不是外形体积明显变小。",
    )
    add_item(
        doc,
        "4. 教师示范，明确制作方法：",
        "教师边操作边讲解浮沉子的结构、装配顺序和调试原则，强调瓶口朝下、保留气室、配重逐次微调。学生对照步骤检查材料，明确操作要求后再开始制作。",
    )
    add_item(
        doc,
        "5. 小组制作，测试调试：",
        "学生以小组为单位完成装配，依次进行漂浮测试、下沉测试和悬浮挑战，并在记录单上写下每一次调整及其结果。教师巡视时不直接给出答案，而是引导学生根据现象判断配重、气室或密封方面的问题。",
    )
    add_item(
        doc,
        "6. 展示交流，解释机制：",
        "各小组展示作品运行效果，说明自己的制作方法、调试过程和成功经验。教师引导学生用“重力与浮力关系”解释上浮、下沉和悬浮三种状态。",
    )
    add_item(
        doc,
        "7. 拓展迁移，总结提升：",
        "联系潜水艇、鱼鳔等实例，引导学生比较不同浮沉控制方式的异同，体会物理知识在工程技术和生命现象中的应用价值。",
    )

    add_heading(doc, "三、浮沉子具体制作流程", 1)
    add_para(
        doc,
        "制作材料：带盖透明塑料瓶1个、装满清水的水槽或烧杯、能留存少量空气的小塑料滴管或开口小药瓶1个、回形针若干、小螺母或橡皮泥少量、抹布。教师课前应检查大瓶是否能密封，浮沉子材料是否清洁无破损。为保证操作安全，不使用玻璃小瓶和尖锐金属件。",
    )
    add_item(
        doc,
        "第一步，认识结构：",
        "教师举起小滴管或小药瓶，请学生找到开口和内部气室。明确制作时开口必须朝下，内部必须留有空气，不能把小容器完全灌满水。",
    )
    add_item(
        doc,
        "第二步，安装配重：",
        "在小滴管或小药瓶下端安装1个回形针，或粘上一小粒橡皮泥，使开口始终朝下。教师提醒学生先少加配重，后续只能一小点一小点地增加，不能一次加得过多。",
    )
    add_item(
        doc,
        "第三步，进行初次漂浮测试：",
        "把浮沉子开口朝下轻轻放入水中。合格的初始状态是浮沉子竖直漂浮，顶部刚好露出水面或接近水面。若大部分露出水面，说明配重过轻；若直接沉底，说明配重过重或内部进水过多。",
    )
    add_item(
        doc,
        "第四步，微调初始状态：",
        "配重过轻时，每次增加一个小回形针或极少量橡皮泥，再重新测试；配重过重时，减少配重或倒出少量内部水。每次只改变一个因素，并把“调整内容—观察现象—测试结果”记录下来。",
    )
    add_item(
        doc,
        "第五步，装入大瓶：",
        "将调好的浮沉子放入装满水的透明塑料瓶中，水面尽量接近瓶口，减少瓶内多余空气。旋紧瓶盖后倒置检查，确认瓶盖处不漏水、不漏气。",
    )
    add_item(
        doc,
        "第六步，测试上浮和下沉：",
        "双手缓慢挤压瓶身，观察浮沉子是否下沉；逐渐松手，观察它是否上浮。教师提醒学生不要猛捏瓶身，要让变化过程清楚可见，并观察挤压时浮沉子内部水面和气室的变化。",
    )
    add_item(
        doc,
        "第七步，挑战悬浮：",
        "先轻轻挤压，使浮沉子缓慢下沉；当它到达瓶身中部时逐渐减小压力，尝试让重力与浮力达到平衡。若总是迅速上浮，说明配重略轻或初始气室偏大；若轻压就快速沉底，说明配重略重或初始气室偏小。",
    )
    add_item(
        doc,
        "第八步，故障排查：",
        "不能下沉时，先检查大瓶是否漏气，再检查浮沉子是否过轻、气室是否过大；不能上浮时，检查配重是否过重、内部是否进水过多；姿态倾斜时，调整配重位置，使重心位于开口下方；状态忽快忽慢时，检查瓶盖密封和浮沉子是否卡在瓶壁。",
    )
    add_item(
        doc,
        "第九步，记录与展示：",
        "作品成功后，记录最终使用的配重数量、最灵敏的挤压力度和实现悬浮的方法。展示时由一名学生操作，一名学生指认气室变化，一名学生用重力和浮力关系解释现象。",
    )

    add_heading(doc, "四、教学活动安排（40分钟）", 1)
    add_item(doc, "导入与任务发布（3分钟）：", "演示现象，明确制作任务。")
    add_item(doc, "原理探究与问题讨论（8分钟）：", "分析浮沉条件和压强传递，突破错误前概念。")
    add_item(doc, "小组制作与调试改进（18分钟）：", "完成浮沉子制作，记录调试过程。")
    add_item(doc, "成果展示与交流评价（7分钟）：", "展示作品，说明原理，开展同伴评价。")
    add_item(doc, "迁移应用与课堂总结（4分钟）：", "联系潜水艇和鱼鳔，总结核心收获。")

    add_heading(doc, "五、评价设计", 1)
    add_para(
        doc,
        "本节课采用过程性评价与结果性评价相结合的方式。过程性评价关注学生是否积极参与小组合作、是否能够记录调试过程、是否能根据失败现象主动改进方案；结果性评价关注作品能否实现上浮、下沉和悬浮，操作是否稳定灵敏；表达性评价关注学生能否用重力和浮力的关系准确解释浮沉机制。通过教师评价、同伴互评和学生自评，使评价真正服务于学习过程。",
    )
    add_heading(doc, "学生学习评价表", 2)
    add_para(doc, "小组：__________　姓名：__________　日期：__________")
    add_evaluation_table(doc)
    add_para(doc, "本组最成功的一次调整：____________________________________________________________")
    add_para(doc, "遇到的主要问题及解决办法：________________________________________________________")
    add_para(doc, "我还想继续改进的地方：____________________________________________________________")

    add_heading(doc, "六、演讲内容", 1)
    add_para(
        doc,
        "接下来是本节课的教学设计。本节课我以“制作浮沉子”为核心任务，采用项目式学习的方式展开。课堂一开始，我会先演示一个能够随手控制上浮、下沉和悬浮的浮沉子，让学生直观感受到这个小装置的神奇之处，并顺势提出本节课的挑战任务：能不能自己制作一个听指挥的小潜水员？",
    )
    add_para(
        doc,
        "在学生产生兴趣之后，我会引导他们观察现象并提出问题。比如，为什么挤压瓶身后浮沉子会下沉？松手后为什么又会上浮？这时学生往往会产生不同猜想，有的会认为是体积变小导致浮力变小，有的会注意到水进入了小瓶内部。这里正是本节课概念突破的关键。",
    )
    add_para(
        doc,
        "因此，在原理探究环节，我会引导学生从受力角度分析浮沉条件，明确物体的浮沉状态取决于重力和浮力的大小关系。通过追问“进入小瓶的水到底改变了谁”，帮助学生认识到，浮沉子下沉的主要原因是水进入内部后自身总重力增大，从而使重力大于浮力。这一过程既突破了学生容易形成的错误前概念，也让他们把浮力知识真正用在解释真实现象上。",
    )
    add_para(
        doc,
        "随后进入小组制作与调试环节。学生根据任务要求选择材料，尝试制作浮沉子，并通过增减配重、调节初始气室大小等方式不断测试和改进。这个环节不是简单的动手操作，而是一个完整的工程迭代过程。学生需要记录失败原因，判断是配重过轻、配重过重、气室不合适，还是存在漏气问题，并据此调整方案。",
    )
    add_para(
        doc,
        "为了让学生真正会做，我会先进行手把手示范。首先让学生辨认浮沉子的开口和气室，明确开口必须朝下、内部必须留有空气；然后从最少的配重开始，把浮沉子放入水中测试。理想状态是它竖直漂浮，顶部刚好接近水面。如果露出水面太多，就一点点增加配重；如果直接沉底，就减少配重或倒出内部的水。调好后再把浮沉子装入盛满水的密封塑料瓶中，通过缓慢挤压和松手测试下沉与上浮，最后挑战在瓶身中部悬浮。整个过程中，每次只改变一个因素，让学生清楚知道是哪一次调整产生了效果。",
    )
    add_para(
        doc,
        "在展示交流环节，各小组展示自己的作品运行效果，并说明调试过程和原理解释。评价不只看作品是否成功，还从作品效果、原理解释、制作调试、合作表现和展示反思五个方面进行。学生先自评，再进行小组互评，教师结合课堂观察给出评价，让学生既看到制作结果，也关注解决问题的过程。最后，我会把浮沉子与潜水艇、鱼鳔进行联系，引导学生比较不同浮沉控制方式的异同，让学生感受到物理知识不仅存在于课堂实验中，也广泛存在于工程技术和生命现象之中。",
    )
    add_para(
        doc,
        "总体来说，这一教学设计希望通过一个看似简单的小制作，让学生经历从现象观察到原理建构、从动手制作到迭代改进、从作品展示到知识迁移的完整学习过程，在理解浮沉条件的同时发展科学思维、工程意识和合作探究能力。",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
