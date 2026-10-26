from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(__file__).resolve().parent / "Part3教学设计PPT文案.docx"


slides = [
    {
        "no": "第19页",
        "title": "第三章 教学设计 [PART3]",
        "body": [
            "制作浮沉子",
            "基于项目式学习的初中物理教学设计",
            "40分钟课堂：在制作与调试中理解浮沉条件",
        ],
        "suggestion": "章节页；保留模板的大标题版式，将“论文内容与讨论”替换为“教学设计”。",
    },
    {
        "no": "第20页",
        "title": "教学设计总览",
        "body": [
            "设计理念：以任务驱动探究，以制作促进理解，以调试深化概念",
            "课堂主线：情境导入 → 原理探究 → 制作调试 → 展示评价 → 迁移总结",
            "核心任务：制作一个能上浮、下沉，并能在水中悬浮的“浮沉子”",
            "学习方式：小组合作、实验观察、工程迭代、证据表达",
        ],
        "suggestion": "建议做成横向流程图，5个环节用箭头连接；每个环节只放关键词。",
    },
    {
        "no": "第21页",
        "title": "教学流程与时间分配",
        "table": [
            ["环节", "时间", "教师活动", "学生活动"],
            ["情境导入", "3分钟", "演示浮沉子“听指挥”沉浮，提出挑战任务", "观察现象，提出疑问"],
            ["原理探究", "8分钟", "引导分析受力变化与压强传递", "猜想原因，比较重力与浮力"],
            ["制作调试", "18分钟", "提供材料，巡视指导，追问失败原因", "小组制作、测试、记录、改进"],
            ["展示评价", "7分钟", "组织成果展示与原理说明", "展示作品，解释沉浮机制"],
            ["迁移总结", "4分钟", "联系潜水艇、鱼鳔，归纳课堂收获", "完成反思，形成迁移认识"],
        ],
        "suggestion": "建议使用四列表格；时间列突出显示，体现40分钟课堂节奏。",
    },
    {
        "no": "第22页",
        "title": "核心活动设计",
        "body": [
            "任务一：看现象——观察浮沉子的上浮、下沉与悬浮",
            "任务二：析原因——判断挤压瓶身后改变的是重力还是浮力",
            "任务三：做作品——选择材料，完成基础浮沉子制作",
            "任务四：调状态——通过增减配重和调节气室实现灵敏控制",
        ],
        "suggestion": "建议做成四个任务卡片或四象限结构；每个任务配一个动词：看、析、做、调。",
    },
    {
        "no": "第23页",
        "title": "问题链与重难点突破",
        "body": [
            "问题1：浮沉子为什么能停在水面附近？",
            "问题2：挤压瓶身后，水是怎样进入小瓶的？",
            "问题3：下沉的关键原因是体积变小，还是总重力变大？",
            "问题4：怎样让浮沉子既不浮起，也不沉底？",
            "问题5：失败时，如何判断是配重、气室还是漏气的问题？",
        ],
        "suggestion": "建议做成递进式问题链，用箭头呈现“现象—原因—证据—调控—排查”。",
    },
    {
        "no": "第24页",
        "title": "评价与拓展",
        "body": [
            "过程评价：是否积极参与制作、记录调试过程、主动分析失败原因",
            "成果评价：作品能否实现上浮、下沉和悬浮，操作是否稳定灵敏",
            "表达评价：能否用“重力与浮力关系”准确解释浮沉机制",
            "合作评价：小组分工是否明确，交流是否充分，改进是否有依据",
            "拓展迁移：比较浮沉子、潜水艇和鱼鳔的浮沉控制方式",
        ],
        "suggestion": "建议分为“评价”和“拓展”两栏；评价可用勾选式要点，拓展配对比图。",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 6 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(9.5)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10.5)


def add_note_table(doc, label, text):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(3.0)
    table.columns[1].width = Cm(13.0)
    row = table.rows[0]
    row.cells[0].width = Cm(3.0)
    row.cells[1].width = Cm(13.0)
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(row.cells[0], "44546A")
    set_cell_shading(row.cells[1], "F3F5F7")
    set_cell_text(row.cells[0], label, bold=True, color="FFFFFF")
    set_cell_text(row.cells[1], text)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color in [
        ("Heading 1", 18, "C00000"),
        ("Heading 2", 14, "44546A"),
    ]:
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def build():
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Part3 教学设计 PPT文案")
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string("C00000")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("适用于《制作浮沉子》中学物理说课 PPT 第19-24页")
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("44546A")

    add_note_table(
        doc,
        "设计定位",
        "40分钟课堂；项目式学习 + 实验探究 + 工程迭代；文案为PPT上屏内容，不含长篇讲稿。",
    )

    for i, slide in enumerate(slides):
        if i:
            doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_heading(f"{slide['no']}  {slide['title']}", level=1)

        doc.add_heading("PPT正文文案", level=2)
        if "table" in slide:
            data = slide["table"]
            table = doc.add_table(rows=len(data), cols=len(data[0]))
            table.style = "Table Grid"
            widths = [Cm(3.0), Cm(2.2), Cm(5.8), Cm(5.8)]
            for r_idx, row in enumerate(data):
                for c_idx, text in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.width = widths[c_idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    if r_idx == 0:
                        set_cell_shading(cell, "44546A")
                        set_cell_text(cell, text, bold=True, color="FFFFFF")
                    else:
                        if c_idx == 1:
                            set_cell_shading(cell, "F8E9E9")
                        set_cell_text(cell, text)
        else:
            for item in slide["body"]:
                add_bullet(doc, item)

        doc.add_paragraph()
        add_note_table(doc, "页面呈现建议", slide["suggestion"])

    doc.save(OUT)


if __name__ == "__main__":
    build()
