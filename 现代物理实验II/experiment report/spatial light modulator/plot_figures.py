from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data.xlsx"
FIG = ROOT / "tex" / "fig"
FIG.mkdir(parents=True, exist_ok=True)


def load_data():
    wb = openpyxl.load_workbook(DATA, data_only=True)
    ws1 = wb["Sheet1"]
    rows1 = list(ws1.iter_rows(min_row=2, values_only=True))
    angle = np.array([r[1] for r in rows1 if r[1] is not None], dtype=float)
    p0 = np.array([r[2] for r in rows1 if r[1] is not None], dtype=float)
    p255 = np.array([r[6] for r in rows1 if r[5] is not None], dtype=float)
    contrast = p255 / p0

    ws2 = wb["Sheet2"]
    rows2 = list(ws2.iter_rows(min_row=2, values_only=True))
    gray = np.array([r[0] for r in rows2 if r[0] is not None], dtype=float)
    power = np.array([r[1] for r in rows2 if r[0] is not None], dtype=float)
    return angle, p0, p255, contrast, gray, power


def style():
    plt.rcParams.update(
        {
            "font.family": ["Hiragino Sans GB", "Arial Unicode MS", "DejaVu Sans"],
            "axes.unicode_minus": False,
            "font.size": 11,
            "axes.linewidth": 1.1,
            "figure.dpi": 160,
            "savefig.dpi": 300,
        }
    )


def save_power_angle(angle, p0, p255):
    def fit_periodic(x, y):
        rad = np.deg2rad(4 * x)
        X = np.column_stack([np.ones_like(x), np.cos(rad), np.sin(rad)])
        beta, *_ = np.linalg.lstsq(X, y, rcond=None)
        return beta

    def eval_periodic(beta, x):
        rad = np.deg2rad(4 * x)
        X = np.column_stack([np.ones_like(x), np.cos(rad), np.sin(rad)])
        return X @ beta

    xfit = np.linspace(0, 350, 600)
    beta0 = fit_periodic(angle, p0)
    beta255 = fit_periodic(angle, p255)

    fig, ax = plt.subplots(figsize=(7.2, 4.6))
    ax.scatter(angle, p0, s=34, marker="o", color="#1f77b4", label="灰度 0 实验点")
    ax.scatter(angle, p255, s=34, marker="s", color="#d95f02", label="灰度 255 实验点")
    ax.plot(xfit, eval_periodic(beta0, xfit), "-", lw=2.0, color="#1f77b4", label="灰度 0 拟合")
    ax.plot(xfit, eval_periodic(beta255, xfit), "-", lw=2.0, color="#d95f02", label="灰度 255 拟合")
    ax.set_xlabel("半波片光轴角度 / (deg)")
    ax.set_ylabel("透过光功率 / μW")
    ax.set_xlim(0, 350)
    ax.grid(True, alpha=0.28)
    ax.legend(frameon=False, ncol=2, loc="upper center", bbox_to_anchor=(0.5, 1.22), fontsize=10)
    fig.tight_layout()
    fig.savefig(FIG / "power_angle.png", bbox_inches="tight")
    plt.close(fig)
    return beta0, beta255


def save_contrast(angle, contrast, beta0, beta255):
    def eval_periodic(beta, x):
        rad = np.deg2rad(4 * x)
        X = np.column_stack([np.ones_like(x), np.cos(rad), np.sin(rad)])
        return X @ beta

    idx = int(np.nanargmax(contrast))
    xfit = np.linspace(0, 350, 1200)
    p0fit = eval_periodic(beta0, xfit)
    p255fit = eval_periodic(beta255, xfit)
    ratio_fit = np.full_like(xfit, np.nan, dtype=float)
    mask = p0fit > 8
    ratio_fit[mask] = p255fit[mask] / p0fit[mask]

    fig, ax = plt.subplots(figsize=(7.2, 4.6))
    ax.scatter(angle, contrast, s=34, color="#b23a48", label="实验点")
    ax.plot(xfit, ratio_fit, "-", color="#1d3557", lw=2.0, label="由功率拟合得到的参考曲线")
    ax.scatter([angle[idx]], [contrast[idx]], s=72, color="#1d3557", zorder=5)
    ax.annotate(
        f"最大值: {contrast[idx]:.1f}\\n角度: {angle[idx]:.0f}°",
        xy=(angle[idx], contrast[idx]),
        xytext=(angle[idx] - 95, contrast[idx] * 0.78),
        arrowprops={"arrowstyle": "->", "lw": 1.0},
        ha="left",
        va="top",
    )
    ax.set_xlabel("半波片光轴角度 / (deg)")
    ax.set_ylabel("对比度 $I_{255}/I_0$")
    ax.set_xlim(0, 350)
    ax.set_ylim(0, max(contrast) * 1.12)
    ax.grid(True, alpha=0.28)
    ax.legend(frameon=False, loc="upper right", fontsize=10)
    fig.tight_layout()
    fig.savefig(FIG / "contrast_angle.png", bbox_inches="tight")
    plt.close(fig)
    return angle[idx], contrast[idx]


def save_gray_response(gray, power):
    coeff = np.polyfit(gray, power, 2)
    xfit = np.linspace(gray.min(), gray.max(), 300)
    yfit = np.polyval(coeff, xfit)
    residual = power - np.polyval(coeff, gray)
    ss_res = float(np.sum(residual**2))
    ss_tot = float(np.sum((power - np.mean(power)) ** 2))
    r2 = 1 - ss_res / ss_tot

    fig, ax = plt.subplots(figsize=(7.2, 4.6))
    ax.scatter(gray, power, s=38, marker="o", color="#1f77b4", label="实验点")
    ax.plot(xfit, yfit, "-", color="#b23a48", lw=2.0, label=f"二次拟合, $R^2={r2:.4f}$")
    ax.set_xlabel("输入图像灰度")
    ax.set_ylabel("输出光功率 / μW")
    ax.set_xlim(0, 255)
    ax.grid(True, alpha=0.28)
    ax.legend(frameon=False, loc="upper left", fontsize=10)
    fig.tight_layout()
    fig.savefig(FIG / "gray_response.png", bbox_inches="tight")
    plt.close(fig)
    return coeff, r2


def main():
    style()
    angle, p0, p255, contrast, gray, power = load_data()
    beta0, beta255 = save_power_angle(angle, p0, p255)
    best_angle, best_contrast = save_contrast(angle, contrast, beta0, beta255)
    coeff, r2 = save_gray_response(gray, power)
    print(f"best_angle={best_angle:.0f}")
    print(f"best_contrast={best_contrast:.6g}")
    print("quadratic_coeff=", " ".join(f"{v:.10g}" for v in coeff))
    print(f"gray_response_r2={r2:.6g}")


if __name__ == "__main__":
    main()
